#!/usr/bin/env python3
# /// script
# requires-python = ">=3.11"
# dependencies = [
#   "python-pptx==1.0.2",
#   "PyYAML>=6.0",
#   "qrcode[pil]>=7.4",
# ]
# ///
"""Build the Kirchhoff AI Builder Cup 2026 submission deck.

    uv run --script docs/submission/deck/build_deck.py            # build the deck
    uv run --script docs/submission/deck/build_deck.py --strict   # also fail while any slot is open
    uv run --script docs/submission/deck/build_deck.py --pdf      # also export a PDF (needs LibreOffice)

Inputs (all relative to this folder):
    ../deck-content.md                     reviewed slide copy, speaker notes, SHOT and TBD slots (binding)
    template/ai-builder-cup-template.pptx  the organisers' 16-slide template (slide 2 is deleted)
    shots/<name>.png|jpg                   screenshots; a missing file renders as a visible placeholder
    fill.yaml                              team-only values; an empty value renders as a highlighted [TBD]
    assets/icons/*.png                     Lucide icons (ISC licence), rasterised at 256 px

Output: Kirchhoff-AI-Builder-Cup.pptx (15 slides) and shots/README.md (the expected screenshot list).

The build is deterministic: the same inputs give a byte-identical .pptx.
"""

from __future__ import annotations

import argparse
import io
import re
import shutil
import subprocess
import sys
import zipfile
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path

import qrcode
import yaml
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.oxml.xmlchemy import OxmlElement
from pptx.text.text import _Paragraph
from pptx.util import Emu, Pt

HERE = Path(__file__).resolve().parent
CONTENT_MD = HERE.parent / "deck-content.md"
TEMPLATE = HERE / "template" / "ai-builder-cup-template.pptx"
SHOTS_DIR = HERE / "shots"
ICONS_DIR = HERE / "assets" / "icons"
FILL_YAML = HERE / "fill.yaml"
OUT_PPTX = HERE / "Kirchhoff-AI-Builder-Cup.pptx"
PATHS = {"shots": SHOTS_DIR}       # overridable from the command line

# Corrections applied to every piece of content text before it reaches a slide or the notes.
# P2-R11: the agents run in ADK agent mode by default, not as ADK workflows.
CORRECTIONS = [
    ("ADK 2.x workflows run the agents on", "ADK 2.x agents (LlmAgent + function tools) run on"),
    ("ADK 2.x workflows", "ADK 2.x agents (LlmAgent + function tools)"),
]

# --------------------------------------------------------------------------------------------
# Design tokens. The template is black and white with Google Sans Flex; colour is semantic only.
# --------------------------------------------------------------------------------------------
INK, INK2, MUTED, FAINT = "202124", "3C4043", "5F6368", "80868B"
LINE, LINE2, SURF, SURF2, WHITE = "DADCE0", "E8EAED", "F8F9FA", "F1F3F4", "FFFFFF"
BLUE, BLUE_T = "1A73E8", "E8F0FE"          # Gemini agents
GREEN, GREEN_T = "188038", "E6F4EA"        # checks that pass
RED, RED_T = "D93025", "FCE8E6"            # the no-LLM boundary
PURPLE, PURPLE_T = "9334E6", "F3E8FD"      # ERP / CRM sandbox
SHOT_LINE, SHOT_FILL, SHOT_TEXT = "C2410C", "FFF4EC", "9A3412"   # screenshot placeholders
TBD_HL, TBD_TEXT = "FFE58F", "5C3D00"      # team-fill placeholders
SPEC_HL, SPEC_TEXT = "FDEBD0", "8A4B00"    # "(spec)" markers

SANS = {400: "Google Sans Flex", 500: "Google Sans Flex Medium", 600: "Google Sans Flex SemiBold"}
MONO = "Roboto Mono"

SLIDE_W, SLIDE_H = 10.0, 5.625
L, R = 0.39, 9.61            # content edges, aligned with the template heading text
CW = R - L
TOP = 1.24                   # first content line under the template heading
FOOT_Y = 5.25                # footer small print

# Advance widths of Google Sans Flex (units per em 2000) for ASCII 32..126 at weights 400 and 600,
# used to estimate wrapping so the build can warn about (and shrink) text that would overflow.
_W400 = [int(v) for v in """450 656 646 1282 1204 1660 1256 360 656 656 854 1116 476 675 476 600 1290 735 1052 1070
 1186 1126 1114 991 1102 1114 476 476 1116 1116 1116 956 1768 1340 1211 1472 1406 1102 1062 1578 1396 490 1060 1252 1018
 1732 1416 1648 1152 1648 1172 1101 1076 1306 1262 1888 1272 1190 1144 761 600 761 930 1125 900 1070 1204 1085 1204 1126
 758 1194 1122 452 452 1010 422 1750 1122 1192 1204 1204 750 943 730 1122 1028 1554 973 1026 963 788 474 788 1116""".split()]
_W600 = [int(v) for v in """443 702 788 1330 1228 1706 1354 462 707 707 907 1164 533 665 533 716 1311 811 1109 1117
 1215 1159 1164 1029 1155 1164 533 533 1164 1164 1164 1062 1779 1385 1268 1440 1434 1161 1102 1583 1463 573 1110 1340 1082
 1843 1464 1635 1228 1635 1268 1158 1138 1355 1337 1970 1356 1291 1181 806 750 806 995 1108 887 1153 1265 1124 1262 1160
 772 1257 1204 530 530 1126 510 1853 1204 1225 1265 1262 832 1000 796 1204 1102 1607 1054 1093 1001 801 575 801 1164""".split()]
_WX = {"·": (476, 533), "–": (1125, 1108), "—": (1800, 1773), "’": (476, 533), "“": (792, 935), "”": (792, 935),
       "…": (1428, 1528), "§": (1092, 1168), "×": (1116, 1164), "≤": (1116, 1164), "≥": (1116, 1164),
       "²": (684, 761), "¹": (501, 584), "≈": (1116, 1164)}
LINE_HEIGHT = 1.25           # (ascender + descender) / em for Google Sans Flex

WARNINGS: list[str] = []


def I(inches: float) -> Emu:
    return Emu(int(round(inches * 914400)))


def rgb(hex_: str) -> RGBColor:
    return RGBColor.from_string(hex_)


def warn(msg: str) -> None:
    WARNINGS.append(msg)


# --------------------------------------------------------------------------------------------
# Content: parse deck-content.md into per-slide sections.
# --------------------------------------------------------------------------------------------
@dataclass
class Section:
    number: int
    title: str
    copy: str
    notes: str
    visual: str

    def copy_lines(self) -> list[str]:
        return [ln.rstrip() for ln in self.copy.splitlines() if ln.strip()]

    def shots(self) -> list[str]:
        return re.findall(r"\[SHOT: (.+?)\]", self.visual)

    def quoted(self, label: str) -> str:
        """Text after `label:` in the Visual block, without surrounding quotes."""
        m = re.search(re.escape(label) + r":\s*(.+?)\s*$", self.visual, re.M)
        if not m:
            raise SystemExit(f"slide {self.number}: no '{label}:' line in the Visual block")
        text = m.group(1).strip()
        if text.startswith('"') and text.endswith('"'):
            text = text[1:-1]
        return text

    def mermaid(self) -> str:
        m = re.search(r"```mermaid\n(.+?)```", self.visual, re.S)
        if not m:
            raise SystemExit(f"slide {self.number}: no mermaid block")
        return m.group(1)


def correct(text: str) -> str:
    for old, new in CORRECTIONS:
        text = text.replace(old, new)
    return text


def parse_content(md: str) -> tuple[dict[int, Section], dict[int, str]]:
    md = correct(md)
    sections: dict[int, Section] = {}
    parts = re.split(r"^## Slide (\d+) · (.+)$", md, flags=re.M)
    for i in range(1, len(parts), 3):
        n, title, body = int(parts[i]), parts[i + 1].strip(), parts[i + 2]
        body = body.split("\n## ")[0]

        def block(name: str, nxt: str | None) -> str:
            pat = r"\*\*" + name + r"\*\*\n(.*?)" + (r"\n\*\*" + nxt + r"\*\*" if nxt else r"\Z")
            m = re.search(pat, body, re.S)
            return (m.group(1) if m else "").strip().removesuffix("---").strip()

        sections[n] = Section(n, title, block("On-slide copy", "Speaker notes"),
                              block("Speaker notes", "Visual"), block("Visual", None))
    sources: dict[int, str] = {}
    m = re.search(r"^## Sources.*?$(.+?)^---", md, re.S | re.M)
    if m:
        for sm in re.finditer(r"^(\d+)\. (.+)$", m.group(1), re.M):
            sources[int(sm.group(1))] = sm.group(2).strip()
    missing = [n for n in range(1, 17) if n not in sections]
    if missing:
        raise SystemExit(f"deck-content.md: missing slide sections {missing}")
    return sections, sources


# --------------------------------------------------------------------------------------------
# Team-fill values, [TBD] slots and (spec) markers.
# --------------------------------------------------------------------------------------------
# [TBD: ...] slots in order of appearance, per slide and part, mapped to fill.yaml keys.
TBD_KEYS = {
    (10, "copy"): ["metrics.usd_per_run"],
    (10, "notes"): ["metrics.tokens_and_usd_per_run"],
    (12, "notes"): ["metrics.upload_to_clean_median", "metrics.intake_accuracy"],
    (14, "copy"): ["links.github_repo", "links.demo_video"],
    (15, "copy"): ["team.mehul_role_background", "team.jayadratha_role_background"],
    (15, "notes"): ["team.notes_bios"],
}
# On-slide "(spec)" markers in order of appearance, mapped to fill.yaml spec_verified keys.
SPEC_KEYS = {
    4: ["s04_power_topologies"],
    6: ["s06_phases_gates"],
    10: ["s10_warm_instance", "s10_budget_alerts"],
}


@dataclass
class OpenItem:
    kind: str      # SHOT, TBD, SPEC, PHOTO
    slide: int     # template slide number (as in deck-content.md)
    key: str
    text: str


class Fill:
    def __init__(self, path: Path):
        self.base = path.parent
        self.data = yaml.safe_load(path.read_text(encoding="utf-8")) if path.exists() else {}
        self.open: list[OpenItem] = []
        self._counters: dict[tuple[int, str], int] = {}
        self._spec_counters: dict[int, int] = {}

    def get(self, dotted: str) -> str:
        node = self.data or {}
        for part in dotted.split("."):
            node = (node or {}).get(part) if isinstance(node, dict) else None
        if node is None or node is False:
            return ""
        text = str(node).strip()
        if re.search(r"[\w.+-]+@[\w-]+\.[a-z]{2,}", text, re.I) or re.search(r"\+\d[\d\s-]{7,}\d", text):
            raise SystemExit(f"fill.yaml {dotted}: looks like an email or phone number; remove it")
        return text

    def tbd(self, slide: int, part: str, inner: str) -> tuple[str, str]:
        keys = TBD_KEYS.get((slide, part), [])
        idx = self._counters.get((slide, part), 0)
        self._counters[(slide, part)] = idx + 1
        if idx >= len(keys):
            raise SystemExit(f"slide {slide} {part}: unexpected [TBD: {inner}]; add it to TBD_KEYS")
        key = keys[idx]
        value = self.get(key)
        if not value:
            self.open.append(OpenItem("TBD", slide, key, inner))
        return key, value

    def spec(self, slide: int) -> tuple[str, bool]:
        keys = SPEC_KEYS.get(slide, [])
        idx = self._spec_counters.get(slide, 0)
        self._spec_counters[slide] = idx + 1
        if idx >= len(keys):
            raise SystemExit(f"slide {slide}: unexpected (spec) marker; add it to SPEC_KEYS")
        key = keys[idx]
        verified = str(self.get(f"spec_verified.{key}")).lower() in ("true", "yes", "1", "verified")
        if not verified:
            self.open.append(OpenItem("SPEC", slide, f"spec_verified.{key}", "(spec) marker shown on slide"))
        return key, verified

    def check_counts(self) -> None:
        for (slide, part), keys in TBD_KEYS.items():
            seen = self._counters.get((slide, part), 0)
            if seen != len(keys):
                raise SystemExit(f"slide {slide} {part}: expected {len(keys)} [TBD] slots, found {seen}")
        for slide, keys in SPEC_KEYS.items():
            if self._spec_counters.get(slide, 0) != len(keys):
                raise SystemExit(f"slide {slide}: expected {len(keys)} (spec) markers on slide")


# --------------------------------------------------------------------------------------------
# Inline text model.
# --------------------------------------------------------------------------------------------
@dataclass
class Run:
    text: str
    bold: bool = False
    italic: bool = False
    mono: bool = False
    kind: str = "text"            # text | tbd | spec
    color: str | None = None
    weight: int | None = None
    scale: float = 1.0
    link: str | None = None


@dataclass
class Para:
    content: str | list[Run]
    size: float | None = None
    weight: int | None = None
    color: str | None = None
    mono: bool = False
    align: str | None = None      # l | c | r
    before: float = 0.0           # points
    after: float = 0.0
    bullet: bool = False
    spc: float | None = None      # character spacing, points
    caps: bool = False
    lnspc: float | None = None


INLINE = re.compile(
    r"\*\*(?P<b>.+?)\*\*|`(?P<c>[^`]+)`|\[TBD: (?P<t>[^\]]+)\]|(?P<s>\s?\(spec\))|(?<![\w*])\*(?P<i>[^*\s][^*]*?)\*(?![\w*])"
)


class Ctx:
    """Resolves inline markdown for one slide: bold, code, italics, [TBD] slots and (spec) markers."""

    def __init__(self, fill: Fill, slide: int):
        self.fill, self.slide = fill, slide

    def runs(self, md: str, **base) -> list[Run]:
        out: list[Run] = []
        pos = 0
        for m in INLINE.finditer(md):
            if m.start() > pos:
                out.append(Run(md[pos:m.start()], **base))
            if m.group("b") is not None:
                out.extend(self.runs(m.group("b"), **{**base, "bold": True}))
            elif m.group("c") is not None:
                out.append(Run(m.group("c"), **{**base, "mono": True}))
            elif m.group("i") is not None:
                out.append(Run(m.group("i"), **{**base, "italic": True}))
            elif m.group("t") is not None:
                _, value = self.fill.tbd(self.slide, "copy", m.group("t"))
                if value:
                    out.append(Run(value, **base))
                else:
                    out.append(Run(f"[TBD: {m.group('t')}]", **{**base, "kind": "tbd"}))
            elif m.group("s") is not None:
                _, verified = self.fill.spec(self.slide)
                if not verified:
                    out.append(Run(" (spec)", **{**base, "kind": "spec", "scale": 0.78}))
            pos = m.end()
        if pos < len(md):
            out.append(Run(md[pos:], **base))
        return out

    def notes(self, md: str) -> str:
        def tbd(m: re.Match) -> str:
            _, value = self.fill.tbd(self.slide, "notes", m.group(1))
            return value if value else m.group(0)

        text = re.sub(r"\[TBD: ([^\]]+)\]", tbd, md)
        text = re.sub(r"\*\*(.+?)\*\*", r"\1", text)
        text = re.sub(r"`([^`]+)`", r"\1", text)
        text = re.sub(r"(?<![\w*])\*([^*\s][^*]*?)\*(?![\w*])", r"\1", text)
        return text.strip()


def plain(md: str) -> str:
    """Strip inline markdown without resolving slots (for diagram labels and checks)."""
    md = re.sub(r"\*\*(.+?)\*\*", r"\1", md)
    md = re.sub(r"`([^`]+)`", r"\1", md)
    return re.sub(r"(?<![\w*])\*([^*\s][^*]*?)\*(?![\w*])", r"\1", md)


# --------------------------------------------------------------------------------------------
# Text measurement (estimate) and rendering.
# --------------------------------------------------------------------------------------------
def _char_em(ch: str, weight: int, mono: bool) -> float:
    if mono:
        return 0.6
    o = ord(ch)
    if 32 <= o < 127:
        a, b = _W400[o - 32], _W600[o - 32]
    elif ch in _WX:
        a, b = _WX[ch]
    elif ch == " ":
        a, b = 450, 443
    else:
        a, b = 1150, 1200
    v = a if weight <= 400 else b if weight >= 600 else (a + b) / 2
    return v / 2000


def _lines(styled: list[tuple[str, float, int, bool]], width_in: float) -> int:
    lines, line_w = 1, 0.0
    for text, size, weight, mono in styled:
        for piece in [text]:
            for tok in re.split(r"(\s+)", piece):
                if not tok:
                    continue
                w = sum(_char_em(c, weight, mono) for c in tok) * size / 72
                if tok.isspace() and tok != " ":
                    line_w += w
                    continue
                if line_w > 0 and line_w + w > width_in:
                    lines += 1
                    line_w = w
                else:
                    line_w += w
                while line_w > width_in:
                    lines += 1
                    line_w -= width_in
    return lines


def text_width(text: str, size: float, weight: int = 400, mono: bool = False) -> float:
    return sum(_char_em(c, weight, mono) for c in text) * size / 72


def _run_weight(run: Run, para: Para, default: int) -> int:
    if run.weight:
        return run.weight
    if run.bold or run.kind == "tbd":
        return 600
    return para.weight or default


@dataclass
class TextStyle:
    size: float = 12
    weight: int = 400
    color: str = INK
    mono: bool = False
    align: str = "l"
    anchor: str = "t"
    margin: tuple[float, float, float, float] = (0, 0, 0, 0)   # l, t, r, b inches
    lnspc: float = 1.0
    min_scale: float = 0.85


def _resolve(paras, ctx: Ctx | None) -> list[tuple[Para, list[Run]]]:
    if isinstance(paras, (str, Para)):
        paras = [paras]
    out = []
    for p in paras:
        if isinstance(p, str):
            p = Para(p)
        if isinstance(p.content, str):
            runs = ctx.runs(p.content) if ctx else [Run(plain(p.content))]
        else:
            runs = p.content
        out.append((p, runs))
    return out


def _measure(resolved, width_in: float, st: TextStyle, scale: float) -> float:
    total = 0.0
    for para, runs in resolved:
        size = (para.size or st.size) * scale
        styled = [(r.text.upper() if para.caps else r.text, size * r.scale, _run_weight(r, para, st.weight),
                   r.mono or para.mono or st.mono) for r in runs] or [("", size, 400, False)]
        indent = 0.16 if para.bullet else 0.0
        n = _lines(styled, max(0.1, width_in - indent))
        lh = max(s for _, s, _, _ in styled)
        total += n * lh * LINE_HEIGHT * (para.lnspc or st.lnspc) / 72 + (para.before + para.after) * scale / 72
    return total


def _set_highlight(run, hex_: str) -> None:
    rPr = run._r.get_or_add_rPr()
    hl = OxmlElement("a:highlight")
    clr = OxmlElement("a:srgbClr")
    clr.set("val", hex_)
    hl.append(clr)
    for tag in ("a:uLnTx", "a:uLn", "a:uFillTx", "a:uFill", "a:latin", "a:ea", "a:cs", "a:sym", "a:hlinkClick"):
        el = rPr.find(qn(tag))
        if el is not None:
            el.addprevious(hl)
            return
    rPr.append(hl)


def _hlink_uses_text_colour(run) -> None:
    hl = run._r.get_or_add_rPr().find(qn("a:hlinkClick"))
    ext_lst = OxmlElement("a:extLst")
    ext = OxmlElement("a:ext")
    ext.set("uri", "{A12FA001-AC4F-418D-AE19-62706E023703}")
    clr = ext.makeelement("{http://schemas.microsoft.com/office/drawing/2018/hyperlinkcolor}hlinkClr", {"val": "tx"})
    ext.append(clr)
    ext_lst.append(ext)
    hl.append(ext_lst)


def _font_name(run: Run, para: Para, st: TextStyle) -> str:
    if run.mono or para.mono or st.mono:
        return MONO
    w = _run_weight(run, para, st.weight)
    return SANS[600 if w >= 600 else 500 if w >= 500 else 400]


ALIGN = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT}
ANCHOR = {"t": MSO_ANCHOR.TOP, "m": MSO_ANCHOR.MIDDLE, "b": MSO_ANCHOR.BOTTOM}


def render(tf, resolved, st: TextStyle, scale: float) -> None:
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    ml, mt, mr, mb = st.margin
    tf.margin_left, tf.margin_top, tf.margin_right, tf.margin_bottom = I(ml), I(mt), I(mr), I(mb)
    tf.vertical_anchor = ANCHOR[st.anchor]
    txBody = tf._txBody
    for p in list(txBody.findall(qn("a:p")))[1:]:
        txBody.remove(p)
    first = tf.paragraphs[0]
    for r in list(first._p):
        if r.tag != qn("a:pPr"):
            first._p.remove(r)
    for i, (para, runs) in enumerate(resolved):
        p = first if i == 0 else tf.add_paragraph()
        p.alignment = ALIGN[para.align or st.align]
        p.line_spacing = para.lnspc or st.lnspc
        if para.before:
            p.space_before = Pt(para.before * scale)
        if para.after:
            p.space_after = Pt(para.after * scale)
        pPr = p._p.get_or_add_pPr()
        if para.bullet:
            pPr.set("marL", str(int(0.16 * 914400)))
            pPr.set("indent", str(int(-0.16 * 914400)))
            bc = OxmlElement("a:buClr")
            c = OxmlElement("a:srgbClr")
            c.set("val", FAINT)
            bc.append(c)
            pPr.append(bc)
            bf = OxmlElement("a:buFont")
            bf.set("typeface", "Arial")
            pPr.append(bf)
            bu = OxmlElement("a:buChar")
            bu.set("char", "•")
            pPr.append(bu)
        else:
            pPr.set("marL", "0")
            pPr.set("indent", "0")
            pPr.append(OxmlElement("a:buNone"))
        size = (para.size or st.size) * scale
        for run in runs:
            r = p.add_run()
            r.text = run.text.upper() if para.caps else run.text
            f = r.font
            f.size = Pt(round(size * run.scale * 2) / 2)
            f.name = _font_name(run, para, st)
            f.bold = False
            f.italic = run.italic or None
            color = run.color or para.color or st.color
            if run.kind == "tbd":
                color = TBD_TEXT
            elif run.kind == "spec":
                color = SPEC_TEXT
            f.color.rgb = rgb(color)
            rPr = r._r.get_or_add_rPr()
            rPr.set("lang", "en-GB")
            if para.spc is not None:
                rPr.set("spc", str(int(para.spc * 100)))
            if run.kind == "tbd":
                _set_highlight(r, TBD_HL)
            elif run.kind == "spec":
                _set_highlight(r, SPEC_HL)
            if run.link:
                r.hyperlink.address = run.link
                _hlink_uses_text_colour(r)


def fit_render(tf, paras, st: TextStyle, w: float, h: float, ctx: Ctx | None, where: str) -> float:
    resolved = _resolve(paras, ctx)
    ml, mt, mr, mb = st.margin
    avail_w, avail_h = w - ml - mr, h - mt - mb
    scale = 1.0
    while _measure(resolved, avail_w, st, scale) > avail_h + 0.005 and scale > st.min_scale:
        scale = round(scale - 0.025, 3)
    if _measure(resolved, avail_w, st, scale) > avail_h + 0.005:
        warn(f"{where}: text may overflow its box ({_measure(resolved, avail_w, st, scale):.2f} in > {avail_h:.2f} in)")
    elif scale < 0.9:
        warn(f"{where}: text shrunk to {scale:.0%} to fit")
    render(tf, resolved, st, scale)
    return scale


def text(slide, x, y, w, h, paras, ctx: Ctx | None = None, name: str = "text", **style) -> object:
    st = TextStyle(**style)
    tb = slide.shapes.add_textbox(I(x), I(y), I(w), I(h))
    tb.name = name
    fit_render(tb.text_frame, paras, st, w, h, ctx, f"{name}")
    return tb


# --------------------------------------------------------------------------------------------
# Shapes, icons, images, placeholders.
# --------------------------------------------------------------------------------------------
def _nostyle(shape) -> None:
    st = shape._element.find(qn("p:style"))
    if st is not None:
        shape._element.remove(st)


def box(slide, x, y, w, h, *, fill=None, line=None, lw=0.75, dash=None, radius=None, shape=None, name=None):
    kind = shape or (MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE)
    s = slide.shapes.add_shape(kind, I(x), I(y), I(w), I(h))
    _nostyle(s)
    if fill:
        s.fill.solid()
        s.fill.fore_color.rgb = rgb(fill)
    else:
        s.fill.background()
    if line:
        s.line.color.rgb = rgb(line)
        s.line.width = Pt(lw)
        if dash:
            s.line.dash_style = dash
    else:
        s.line.fill.background()
    if radius and kind == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = min(0.5, radius / min(w, h))
    if name:
        s.name = name
    tf = s.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    return s


def box_text(shape, paras, ctx: Ctx | None = None, **style) -> None:
    st = TextStyle(**style)
    w, h = shape.width / 914400, shape.height / 914400
    fit_render(shape.text_frame, paras, st, w, h, ctx, shape.name)


def connector(slide, x1, y1, x2, y2, *, kind="straight", color=MUTED, lw=1.0, dash=None, head=False, tail=True,
              begin=None, end=None, name="arrow"):
    ctype = MSO_CONNECTOR.ELBOW if kind in ("elbow", "hv") else MSO_CONNECTOR.STRAIGHT
    c = slide.shapes.add_connector(ctype, I(x1), I(y1), I(x2), I(y2))
    _nostyle(c)
    if kind == "hv":            # one bend: horizontal out of the start, vertical into the end
        c._element.spPr.find(qn("a:prstGeom")).set("prst", "bentConnector2")
    if begin is not None:
        c.begin_connect(*begin)
    if end is not None:
        c.end_connect(*end)
    c.line.color.rgb = rgb(color)
    c.line.width = Pt(lw)
    if dash:
        c.line.dash_style = dash
    ln = c.line._get_or_add_ln()
    for tag, on in (("a:headEnd", head), ("a:tailEnd", tail)):
        if on:
            el = OxmlElement(tag)
            el.set("type", "triangle")
            el.set("w", "med")
            el.set("len", "med")
            ln.append(el)
    c.name = name
    return c


def rule(slide, x1, y1, x2, y2, color=LINE, lw=0.75, name="rule"):
    return connector(slide, x1, y1, x2, y2, color=color, lw=lw, tail=False, name=name)


def balance(label: str, width_in: float, size: float, weight: int = 400) -> list[str]:
    """Split a label into at most three lines of even width, breaking only at spaces."""
    if text_width(label, size, weight) <= width_in:
        return [label]
    words = label.split(" ")
    best, best_w = [label], 1e9
    for i in range(1, len(words)):
        two = [" ".join(words[:i]), " ".join(words[i:])]
        w = max(text_width(t, size, weight) for t in two)
        if w < best_w:
            best, best_w = two, w
    if best_w <= width_in:
        return best
    for i in range(1, len(words)):
        for j in range(i + 1, len(words)):
            three = [" ".join(words[:i]), " ".join(words[i:j]), " ".join(words[j:])]
            w = max(text_width(t, size, weight) for t in three)
            if w < best_w:
                best, best_w = three, w
    return best


def icon(slide, name: str, x, y, size, color="ink"):
    path = ICONS_DIR / f"{name}-{color}.png"
    if not path.exists():
        raise SystemExit(f"missing icon {path.name}")
    pic = slide.shapes.add_picture(str(path), I(x), I(y), I(size), I(size))
    pic.name = f"icon {name}"
    pic._element.nvPicPr.cNvPr.set("descr", "")     # decorative
    return pic


def icon_badge(slide, name: str, x, y, d, color="ink", bg=SURF2):
    box(slide, x, y, d, d, fill=bg, shape=MSO_SHAPE.OVAL, name=f"badge {name}")
    pad = d * 0.24
    icon(slide, name, x + pad, y + pad, d - 2 * pad, color)


def number_badge(slide, n: str, x, y, d=0.22, fill=INK, fg=WHITE, size=8):
    b = box(slide, x, y, d, d, fill=fill, shape=MSO_SHAPE.OVAL, name=f"number {n}")
    box_text(b, Para(n, align="c"), size=size, weight=600, color=fg, anchor="m")
    return b


def kicker(slide, x, y, w, label: str, color=MUTED, size=7.0, name="kicker"):
    return text(slide, x, y, w, 0.2, Para(label, caps=True, spc=0.8, mono=True), size=size, color=color,
                weight=500, name=name)


def footer(slide, line: str, ctx: Ctx | None = None):
    text(slide, L, FOOT_Y, CW, 0.2, Para(line), ctx, name="footer small print", size=7.5, color=FAINT,
         anchor="b", min_scale=0.9)


def find_image(stem: str) -> Path | None:
    for ext in (".png", ".jpg", ".jpeg", ".PNG", ".JPG", ".JPEG"):
        p = PATHS["shots"] / f"{stem}{ext}"
        if p.exists():
            return p
    return None


def picture(slide, path: Path, x, y, w, h, mode="contain", border=True, name="screenshot", valign="m"):
    with Image.open(path) as im:
        iw, ih = im.size
    box_ar, img_ar = w / h, iw / ih
    if mode == "cover":
        pic = slide.shapes.add_picture(str(path), I(x), I(y), I(w), I(h))
        if img_ar > box_ar:
            cut = (1 - box_ar / img_ar) / 2
            pic.crop_left = pic.crop_right = cut
        else:
            cut = (1 - img_ar / box_ar) / 2
            pic.crop_top = pic.crop_bottom = cut
    else:
        if img_ar > box_ar:
            pw, ph = w, w / img_ar
        else:
            pw, ph = h * img_ar, h
        top = y if valign == "t" else y + (h - ph) / 2
        pic = slide.shapes.add_picture(str(path), I(x + (w - pw) / 2), I(top), I(pw), I(ph))
    if border:
        pic.line.color.rgb = rgb(LINE)
        pic.line.width = Pt(0.75)
    pic.name = name
    return pic


class Shots:
    """Screenshot slots: each [SHOT] in deck-content.md maps to a named file in shots/."""

    FILES = {
        3: ["s03-workspace-signed"],
        5: ["s05-1-intake-quotes", "s05-2-hall-3d-tiers", "s05-3-compliance-matrix", "s05-4-option-r15-swap",
            "s05-5-impact-chain", "s05-6-copilot-voice", "s05-7-review-stamp", "s05-8-integrations",
            "s05-9-solar-field"],
        7: ["s07-finding-drawer-r05"],
        11: ["s11-1-projects-dashboard", "s11-2-hall-b-live-run", "s11-3-finding-drawer-r11",
             "s11-4-options-r15-swap", "s11-5-what-if-200kw", "s11-6-review-approved"],
        12: ["s12-compliance-matrix-first-pass"],
        13: ["s13-grid-pack", "s13-solar-pack"],
    }
    SHAPE = {3: "16:9 full screen", 5: "tight crop of the named element, about 3:2 (centre-cropped to fit)",
             7: "portrait crop of the drawer, trimmed to its content (about 0.55 wide to tall)", 11: "16:9 full screen", 12: "16:9 full screen",
             13: "16:9 full screen"}

    def __init__(self, sections: dict[int, Section], fill: Fill):
        self.fill = fill
        self.slots: dict[int, list[tuple[str, str]]] = {}
        for n, stems in self.FILES.items():
            texts = sections[n].shots()
            if len(texts) != len(stems):
                raise SystemExit(f"slide {n}: {len(texts)} [SHOT] slots in deck-content.md, "
                                 f"{len(stems)} file names in Shots.FILES; update the map")
            self.slots[n] = list(zip(stems, texts))
        extra = [n for n, s in sections.items() if s.shots() and n not in self.FILES]
        if extra:
            raise SystemExit(f"slides {extra} have [SHOT] slots without file names")

    def place(self, slide, n: int, i: int, x, y, w, h, mode="contain", compact=False, valign="m"):
        stem, desc = self.slots[n][i]
        path = find_image(stem)
        if path:
            pic = picture(slide, path, x, y, w, h, mode=mode, name=f"screenshot {stem}", valign=valign)
            pic._element.nvPicPr.cNvPr.set("descr", desc)
            return pic
        self.fill.open.append(OpenItem("SHOT", n, f"shots/{stem}.png", desc))
        return placeholder(slide, x, y, w, h, f"shots/{stem}.png", desc, compact=compact)

    def readme(self) -> str:
        lines = [
            "# Deck screenshots",
            "",
            "`build_deck.py` places each file below on its slide. A missing file renders as an orange dashed",
            "placeholder that shows the SHOT text. Use `.png` (or `.jpg`) with exactly these names, capture",
            "every value from the recorded golden run, and keep the app's fictional-data banner visible.",
            "Full-screen captures at 1920 × 1080 (16:9) fit best; slide 5 tiles are centre-cropped to fit.",
            "",
            "This file is regenerated by every build from `../../deck-content.md`.",
            "",
            "`s05-6-copilot-voice.png` (a real Gemini Live exchange) and `s11-2-hall-b-live-run.png` (a frame of a live run) are captured by hand on the",
            "deployed site: `e2e/shots/capture.spec.ts` skips them, so a regeneration never overwrites them.",
            "",
            "| File | Template slide | Frame | SHOT (from deck-content.md) |",
            "|---|---|---|---|",
        ]
        for n, slots in self.slots.items():
            for stem, desc in slots:
                lines.append(f"| `{stem}.png` | {n} | {self.SHAPE[n]} | {desc.replace('|', '/')} |")
        lines += ["", "Team headshots for slide 15 are set in `../fill.yaml` (`photos.mehul`, `photos.jayadratha`).", ""]
        return "\n".join(lines)


def placeholder(slide, x, y, w, h, label: str, desc: str, compact=False):
    b = box(slide, x, y, w, h, fill=SHOT_FILL, line=SHOT_LINE, lw=1.0, dash=MSO_LINE.DASH, radius=0.06,
            name=f"SHOT placeholder {label}")
    pad = 0.07 if compact else 0.12
    isz = 0.16 if compact else 0.22
    icon(slide, "camera", x + pad, y + pad, isz, "orange")
    head_size = 6.0 if compact else 7.0
    text(slide, x + pad + isz + 0.06, y + pad - 0.01, w - 2 * pad - isz - 0.06, isz + 0.04,
         Para(("SCREENSHOT  " if not compact else "") + label, mono=True, weight=600),
         size=head_size, color=SHOT_TEXT, anchor="m", name="SHOT label", min_scale=0.7)
    body_y = y + pad + isz + 0.05
    text(slide, x + pad, body_y, w - 2 * pad, h - (body_y - y) - pad, Para(desc), size=6.5 if compact else 8.0,
         color=SHOT_TEXT, name="SHOT text", min_scale=0.75)
    return b


def card(slide, x, y, w, h, fill=SURF, line=LINE2, radius=0.08, name="card"):
    return box(slide, x, y, w, h, fill=fill, line=line, lw=0.75, radius=radius, name=name)


def qr_png(url: str) -> io.BytesIO:
    q = qrcode.QRCode(error_correction=qrcode.constants.ERROR_CORRECT_M, box_size=12, border=1)
    q.add_data(url)
    q.make(fit=True)
    img = q.make_image(fill_color="#202124", back_color="white").convert("RGB")
    buf = io.BytesIO()
    img.save(buf, format="PNG", optimize=False)
    buf.seek(0)
    return buf


# --------------------------------------------------------------------------------------------
# Template helpers.
# --------------------------------------------------------------------------------------------
def shape_with_text(slide, needle: str):
    for sh in slide.shapes:
        if sh.has_text_frame and needle in sh.text_frame.text:
            return sh
    raise SystemExit(f"template: no shape containing {needle!r}")


def keep_heading_only(slide, heading: str) -> None:
    """Reduce a template heading text box to its heading line, in the template's heading style."""
    sh = shape_with_text(slide, heading)
    txBody = sh.text_frame._txBody
    paras = txBody.findall(qn("a:p"))
    for p in paras[1:]:
        txBody.remove(p)
    p0 = paras[0]
    for br in p0.findall(qn("a:br")):
        p0.remove(br)
    for r in p0.findall(qn("a:r")):
        for tag in ("a:latin", "a:ea", "a:cs", "a:sym"):
            el = r.find(qn("a:rPr")).find(qn(tag))
            if el is not None:
                el.set("typeface", SANS[500])
    end = p0.find(qn("a:endParaRPr"))
    if end is not None:
        for tag in ("a:latin", "a:ea", "a:cs", "a:sym"):
            el = end.find(qn(tag))
            if el is not None:
                el.set("typeface", SANS[500])
    sh.height = I(0.615)


def add_heading(slide, heading: str) -> None:
    """Heading for the blank template slide, in the same box and style as the other headings."""
    tb = slide.shapes.add_textbox(Emu(265350), Emu(663825), Emu(8943000), Emu(561900))
    tb.name = "heading"
    tf = tb.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = Emu(91425)
    tf.word_wrap = True
    r = tf.paragraphs[0].add_run()
    r.text = heading
    r.font.size = Pt(16)
    r.font.name = SANS[500]
    r.font.color.rgb = rgb("000000")


def set_notes(slide, notes: str) -> None:
    tf = slide.notes_slide.notes_text_frame
    if tf is None:
        raise SystemExit("template notes slide has no body placeholder")
    tf.text = notes


def delete_slide(prs, index: int) -> None:
    sldIdLst = prs.slides._sldIdLst
    sldId = list(sldIdLst)[index]
    prs.part.drop_rel(sldId.rId)
    sldIdLst.remove(sldId)


# --------------------------------------------------------------------------------------------
# Slides. Each builder takes the template slide, its content section and a context.
# --------------------------------------------------------------------------------------------
@dataclass
class Deck:
    sections: dict[int, Section]
    sources: dict[int, str]
    fill: Fill
    shots: Shots
    tagline: str = ""
    product: str = ""
    by: str = ""


def s01_team(slide, sec: Section, ctx: Ctx, deck: Deck):
    lines = sec.copy_lines()
    vals = {}
    for ln in lines:
        m = re.match(r"^([abc])\. ([^:]+): (.+)$", ln)
        if m:
            vals[m.group(1)] = (m.group(2), m.group(3))
    if set(vals) != {"a", "b", "c"}:
        raise SystemExit("slide 1: expected a./b./c. lines")

    sh = shape_with_text(slide, "Team Details")
    sh.left, sh.top, sh.width, sh.height = I(0.16), I(3.06), I(6.0), I(2.16)
    tf = sh.text_frame
    txBody = tf._txBody
    paras = txBody.findall(qn("a:p"))
    # Keep the heading paragraph and the template's a./b./c. auto-numbered items; drop the empty lines.
    heading, items = paras[0], paras[2:5]
    for p in (paras[1], paras[5]):
        txBody.remove(p)
    for r in heading.findall(".//" + qn("a:rPr")) + heading.findall(qn("a:endParaRPr")):
        r.set("sz", "1800")
    spc = OxmlElement("a:spcPts")
    spc.set("val", "500")
    heading.find(qn("a:pPr")).find(qn("a:spcAft")).clear()
    heading.find(qn("a:pPr")).find(qn("a:spcAft")).append(spc)
    for key, p in zip("abc", items):
        pPr = p.find(qn("a:pPr"))
        pPr.set("marL", str(int(0.5 * 914400)))
        pPr.set("indent", str(int(-0.27 * 914400)))
        pPr.find(qn("a:buSzPts")).set("val", "1250")
        sa = pPr.find(qn("a:spcAft"))
        sa.clear()
        s2 = OxmlElement("a:spcPts")
        s2.set("val", "300")
        sa.append(s2)
        label_r = p.find(qn("a:r"))
        rpr = label_r.find(qn("a:rPr"))
        rpr.set("sz", "1250")
        for tag in ("a:latin", "a:ea", "a:cs", "a:sym"):
            rpr.find(qn(tag)).set("typeface", SANS[500])
        p.find(qn("a:endParaRPr")).set("sz", "1250")
        label, value = vals[key]
        label_r.find(qn("a:t")).text = f"{label}: "
        para = _Paragraph(p, tf)        # value runs go after the template's label run
        for run in ctx.runs(value):
            r = para.add_run()
            r.text = run.text
            r.font.size = Pt(12.5)
            r.font.name = SANS[600] if key in "ab" else SANS[400]
            r.font.color.rgb = rgb(INK if key in "ab" else INK2)
            r._r.get_or_add_rPr().set("lang", "en-GB")

    # Wordmark and tagline lockup beside the team block.
    rule(slide, 6.42, 3.3, 6.42, 5.08, name="divider")
    text(slide, 6.72, 3.2, 2.9, 0.55, Para(deck.product, weight=600), size=30, color=INK, name="wordmark",
         anchor="b")
    text(slide, 6.74, 3.78, 2.9, 0.25, Para(deck.by), size=11, color=MUTED, name="by line")
    dots = [BLUE, GREEN, INK]
    for i, sentence in enumerate(re.findall(r"[^.]+\.", deck.tagline)):
        y = 4.2 + i * 0.29
        box(slide, 6.75, y + 0.085, 0.09, 0.09, fill=dots[i], shape=MSO_SHAPE.OVAL, name="tagline dot")
        text(slide, 6.95, y, 2.7, 0.27, Para(sentence.strip(), weight=500), size=13, color=INK2,
             name="tagline", anchor="m")
    footer(slide, sec.quoted("Footer"), ctx)


def s03_brief(slide, sec: Section, ctx: Ctx, deck: Deck):
    paras = [p.strip() for p in sec.copy.split("\n\n") if p.strip()]
    headline, body = paras[0], paras[1]
    head_lines = [s.strip() for s in re.findall(r"[^.]+\.", headline)]
    text(slide, L, TOP + 0.04, 3.75, 0.72, [Para(s, weight=500) for s in head_lines], ctx, size=18,
         color=INK, name="headline", lnspc=1.0)
    sentences = [s.strip() for s in re.findall(r"[^.]+\.", body)]
    if len(sentences) != 4:
        raise SystemExit("slide 3: expected four sentences in the body copy")
    rows = [("sparkles", "blue"), ("shield-check", "green"), ("package", "purple"), ("signature", "ink")]
    y = 2.22
    for (ic, col), s in zip(rows, sentences):
        icon_badge(slide, ic, L, y, 0.36, col, bg=SURF2)
        text(slide, L + 0.5, y - 0.04, 3.25, 0.46, Para(s), ctx, size=11.5, color=INK2, anchor="m",
             name="brief row")
        y += 0.62
    # Lockup above the shot, then the hero screenshot.
    sx, sw = 4.42, R - 4.42
    sh = sw * 9 / 16
    text(slide, sx, TOP, sw, 0.3, [Para([Run(deck.product, weight=600, color=INK), Run("   "),
                                         Run(deck.tagline, color=MUTED)])], size=11, anchor="m",
         name="lockup")
    deck.shots.place(slide, 3, 0, sx, TOP + 0.4, sw, sh)
    footer(slide, sec.quoted("Footer (small print)"), ctx)


def s04_opportunities(slide, sec: Section, ctx: Ctx, deck: Deck):
    keep_heading_only(slide, "Opportunities")
    qa: list[tuple[str, str]] = []
    q = None
    for ln in sec.copy_lines():
        m = re.match(r"^\*(.+)\*$", ln)
        if m:
            q = m.group(1)
        elif q:
            qa.append((q, ln))
            q = None
    if len(qa) != 3:
        raise SystemExit("slide 4: expected three question/answer pairs")
    widths = [3.9, 2.51, 2.51]
    x = L
    for i, ((question, answer), w) in enumerate(zip(qa, widths)):
        card(slide, x, TOP, w, 1.42, name="question card")
        text(slide, x + 0.16, TOP + 0.1, w - 0.32, 0.34, Para(question), size=9, color=MUTED,
             name="template prompt", anchor="m")
        if i == 0:
            text(slide, x + 0.16, TOP + 0.5, w - 0.32, 0.84, Para(answer), ctx, size=11.5, color=INK,
                 name="answer")
        else:
            text(slide, x + 0.16, TOP + 0.5, w - 0.32, 0.84, [Para(s.strip(), weight=500)
                 for s in re.findall(r"[^.]+\.", answer)], ctx, size=13.5, color=INK, name="answer")
        x += w + 0.15

    # Comparison grid, built from shapes so cells can carry icons.
    rows = [ln for ln in sec.visual.splitlines() if ln.startswith("|")]
    table = [[c.strip() for c in r.strip().strip("|").split("|")] for r in rows if not re.match(r"^\|[-| ]+\|$", r)]
    header, body = table[0], table[1:]
    gy = 2.86
    col_w = [2.24, 1.5] + [(CW - 2.24 - 1.5) / 4] * 4
    row_h, head_h = 0.35, 0.32
    xs = [L]
    for w in col_w[:-1]:
        xs.append(xs[-1] + w)
    box(slide, xs[1] + 0.03, gy - 0.04, col_w[1] - 0.06, head_h + row_h * len(body) + 0.08, fill=BLUE_T,
        radius=0.06, name="Kirchhoff column")
    for j, name in enumerate(header):
        if j == 0:
            continue
        text(slide, xs[j] + 0.1, gy, col_w[j] - 0.2, head_h, Para(name, weight=600 if j == 1 else 500),
             size=8.5, color=BLUE if j == 1 else INK2, anchor="m", name="grid header")
    for i, row in enumerate(body):
        ry = gy + head_h + i * row_h
        box(slide, L, ry, CW, 0.01, fill=LINE2, name="grid rule")
        text(slide, L, ry, col_w[0] - 0.1, row_h, Para(row[0]), size=8.5, color=INK, anchor="m",
             name="grid row label")
        for j in range(1, len(header)):
            cell = row[j]
            cx, cw = xs[j] + 0.1, col_w[j] - 0.2
            if cell.startswith("✓"):
                rest = cell[1:].strip()
                icon(slide, "check", cx, ry + (row_h - 0.16) / 2, 0.16, "green")
                if rest:
                    text(slide, cx + 0.2, ry, cw - 0.2, row_h, Para(rest, weight=500), ctx, size=8.5,
                         color=INK, anchor="m", name="grid cell")
            else:
                shown = not cell.startswith("not shown")
                text(slide, cx, ry, cw, row_h, Para(cell), ctx, size=8.5, color=INK2 if shown else FAINT,
                     anchor="m", name="grid cell")
    box(slide, L, gy + head_h + row_h * len(body), CW, 0.01, fill=LINE2, name="grid rule")
    footer(slide, sec.quoted("Footnote"), ctx)


def s05_features(slide, sec: Section, ctx: Ctx, deck: Deck):
    labels = [ln[2:] for ln in sec.copy_lines() if ln.startswith("- ")]
    if len(labels) != 9:
        raise SystemExit("slide 5: expected nine feature bullets")
    icons = ["file-check", "plug-zap", "shield-check", "truck", "waypoints", "mic", "stamp", "arrow-right-left", "sun"]
    gap = 0.14
    tw, th = (CW - 2 * gap) / 3, (FOOT_Y - 0.12 - TOP - 2 * gap) / 3
    for i, label in enumerate(labels):
        cx, cy = L + (i % 3) * (tw + gap), TOP + (i // 3) * (th + gap)
        card(slide, cx, cy, tw, th, name="feature tile")
        icon_badge(slide, icons[i], cx + 0.12, cy + 0.12, 0.34, bg=WHITE)
        text(slide, cx + 0.12, cy + 0.52, 1.18, th - 0.6, Para(label, weight=500), ctx, size=10, color=INK,
             name="feature label")
        iw = tw - 1.42
        deck.shots.place(slide, 5, i, cx + 1.36, cy + 0.08, iw - 0.02, th - 0.16, mode="cover", compact=True)
    footer(slide, sec.quoted("Footer (small print)"), ctx)


def _mermaid_nodes(src: str) -> dict[str, str]:
    nodes = {}
    for m in re.finditer(r'(\w+)(?:\[\("|\["|\{")(.+?)(?:"\)\]|"\]|"\})', src):
        nodes.setdefault(m.group(1), m.group(2))
    return nodes


def _mermaid_edge_labels(src: str) -> list[str]:
    return re.findall(r'-->\|"(.+?)"\|', src)


def _label(nodes: dict[str, str], key: str) -> list[str]:
    if key not in nodes:
        raise SystemExit(f"mermaid: node {key} not found")
    return [s.strip() for s in nodes[key].split("<br/>")]


def diagram_card(slide, x, y, w, h, title: str, sub: str = "", *, fill=WHITE, line=LINE, color=INK, size=8.5,
                 sub_size=7.5, name="node", shape=None, radius=0.05, lw=0.75, title_weight=600, align="c"):
    b = box(slide, x, y, w, h, fill=fill, line=line, lw=lw, radius=radius if shape is None else None,
            shape=shape, name=name)
    inner = w - 0.12
    paras = [Para(t, weight=title_weight, size=size, color=color, align=align)
             for t in balance(title, inner, size, title_weight)]
    if sub:
        paras += [Para(t, size=sub_size, color=INK2 if color == INK else color, align=align)
                  for t in balance(sub, inner, sub_size)]
    box_text(b, paras, size=size, anchor="m", margin=(0.06, 0.03, 0.06, 0.03), min_scale=0.8)
    return b


def edge_label(slide, x, y, w, h, label: str, align="c", bg=WHITE):
    lines = balance(label, w - 0.02, 7)
    tb = text(slide, x, y, w, h, [Para(t, align=align) for t in lines], size=7, color=MUTED, anchor="m",
              name="edge label", min_scale=0.85)
    if bg:
        tb.fill.solid()
        tb.fill.fore_color.rgb = rgb(bg)
    return tb


def s06_process(slide, sec: Section, ctx: Ctx, deck: Deck):
    paras = [p.strip() for p in sec.copy.split("\n\n") if p.strip()]
    text(slide, L, TOP - 0.04, 5.6, 0.36, Para(paras[0], weight=500), ctx, size=15, color=INK, name="headline",
         anchor="m")
    text(slide, L, TOP + 0.33, CW, 0.26, Para(paras[1]), ctx, size=10, color=INK2, name="subline", anchor="m")

    lanes_m = re.search(r"four lanes: (.+?)\. Show", sec.visual)
    lanes = [plain(s).strip() for s in lanes_m.group(1).split(" · ")]
    phases = re.findall(r"\*\*(P\d): ([^*]+?)\.\*\*", sec.notes)
    if len(lanes) != 4 or len(phases) != 4:
        raise SystemExit("slide 6: expected four lanes and four phases")
    nodes = _mermaid_nodes(sec.mermaid())
    t_in = _label(nodes, "T")[0]
    p1 = _label(nodes, "P1")
    p2 = _label(nodes, "P2")
    p3 = _label(nodes, "P3")
    p4 = _label(nodes, "P4")
    gates = [_label(nodes, g) for g in ("G1", "G2", "G3")]

    top, head_h = 1.86, 0.34
    lane_y0 = top + head_h
    lane_h = (5.2 - lane_y0) / 4
    lw_ = 1.0
    fx = L + lw_ + 0.06
    cols = [0.86, 1.48, 0.45, 1.48, 0.45, 1.48, 0.45, 0]       # input, P1, gate, P2, gate, P3, gate, P4
    cols[-1] = R - fx - sum(cols[:-1])
    cx = [fx]
    for w in cols[:-1]:
        cx.append(cx[-1] + w)
    lane_colors = [INK, BLUE, RED, PURPLE]
    for i, lane in enumerate(lanes):
        y = lane_y0 + i * lane_h
        box(slide, L, y, CW, lane_h, fill=SURF if i % 2 == 0 else WHITE, name="lane band")
        rule(slide, L, y, R, y, name="lane rule")
        lane_lines = balance(lane.replace("(no LLM)", "(no\u00a0LLM)"), lw_ - 0.14, 8, 600)
        text(slide, L + 0.08, y, lw_ - 0.1, lane_h, [Para(t, weight=600) for t in lane_lines], size=8,
             color=lane_colors[i], anchor="m", name="lane label")
    rule(slide, L, lane_y0 + 4 * lane_h, R, lane_y0 + 4 * lane_h, name="lane rule")
    # The deterministic lane: dashed red boundary around its contents, as on the architecture slide.
    box(slide, fx - 0.04, lane_y0 + 2 * lane_h + 0.04, R - fx + 0.02, lane_h - 0.08, line=RED, lw=1.0,
        dash=MSO_LINE.DASH, radius=0.04, name="no-LLM lane boundary")
    # Phase headers over the P columns.
    for k, (pid, pname) in enumerate(phases):
        x, w = cx[1 + 2 * k], cols[1 + 2 * k]
        text(slide, x, top, w, head_h, [Para(pid, mono=True, size=7, color=MUTED, weight=500),
                                         Para(pname[0].upper() + pname[1:], weight=600)], size=8, color=INK,
             align="c", name="phase header", lnspc=0.95)

    def lane_box(lane, col, title, sub="", fill=WHITE, line=LINE, color=INK, h=0.56, inset=(0.07, 0.07)):
        x, w = cx[col] + inset[0], cols[col] - inset[0] - inset[1]
        y = lane_y0 + lane * lane_h + (lane_h - h) / 2
        return diagram_card(slide, x, y, w, h, title, sub, fill=fill, line=line, color=color, size=8,
                            sub_size=7.2, name=f"lane{lane} box")

    tender = lane_box(0, 0, t_in, h=0.46, line=INK2, inset=(0.0, 0.08))
    a1 = lane_box(1, 1, p1[0].removeprefix("P1 "), p1[1], fill=BLUE_T, line=BLUE)
    a2 = lane_box(1, 3, p2[0].removeprefix("P2 ").replace(" ∥ ", "\u00a0∥\u00a0"), p2[1], fill=BLUE_T, line=BLUE,
                  h=0.6)
    e2 = lane_box(2, 3, p2[2][0].upper() + p2[2][1:])
    erp_lead = "ERP lead times"
    if erp_lead not in p2[1]:
        warn("slide 6: 'ERP lead times' no longer in the P2 label")
    r2 = lane_box(3, 3, erp_lead, fill=PURPLE_T, line=PURPLE)
    p3_route, p3_boq = [s.strip() for s in p3[0].removeprefix("P3 ").split(",", 1)]
    checker_title, p3_loop = [s.strip() for s in p3[1].split("),", 1)]
    checker_title += ")"                                   # "Checker (no LLM)", "fix loop"
    a3 = lane_box(1, 5, p3_route, fill=BLUE_T, line=BLUE)
    e3 = lane_box(2, 5, f"{p3_boq} · {checker_title}")
    p4_items = [s.strip().rstrip(",") for s in (p4[0].removeprefix("P4 ") + " " + p4[1]).split(",")]
    p4_items = [s for s in p4_items if s]
    r12 = [s for s in p4_items if s.startswith("R12")]
    rest = [s for s in p4_items if not s.startswith("R12")]
    e4 = lane_box(2, 7, r12[0] if r12 else "R12 re-check", inset=(0.3, 0.07))
    r4 = lane_box(3, 7, " · ".join(rest), fill=PURPLE_T, line=PURPLE, h=0.6, inset=(0.12, 0.07))

    # Gates: diamonds low in the engineer lane, labels above them.
    d = 0.32
    diamonds = []
    for k, g in enumerate(gates):
        gx = cx[2 + 2 * k] + cols[2 + 2 * k] / 2
        gy = lane_y0 + lane_h - d - 0.08
        dm = box(slide, gx - d / 2, gy, d, d, fill=INK, shape=MSO_SHAPE.DIAMOND, name=f"gate {g[0]}")
        diamonds.append(dm)
        label = [Para(g[0], weight=600)] + ([Para(g[1], size=7, color=MUTED)] if len(g) > 1 else [])
        text(slide, gx - 0.6, lane_y0 + 0.03, 1.2, gy - lane_y0 - 0.05, label, size=7.5, color=INK, align="c",
             anchor="b", name="gate label", lnspc=0.92)

    E = lambda s, i: (s, i)  # noqa: E731  (shape, connection site: 0 top, 1 left, 2 bottom, 3 right)
    connector(slide, 0, 0, 1, 1, kind="hv", begin=E(tender, 3), end=E(a1, 0))
    connector(slide, 0, 0, 1, 1, kind="hv", begin=E(a1, 3), end=E(diamonds[0], 2))
    connector(slide, 0, 0, 1, 1, kind="hv", begin=E(diamonds[0], 3), end=E(a2, 0))
    connector(slide, 0, 0, 1, 1, begin=E(a2, 2), end=E(e2, 0))
    connector(slide, 0, 0, 1, 1, begin=E(r2, 0), end=E(e2, 2), color=PURPLE)
    connector(slide, 0, 0, 1, 1, kind="hv", begin=E(e2, 3), end=E(diamonds[1], 2))
    connector(slide, 0, 0, 1, 1, kind="hv", begin=E(diamonds[1], 3), end=E(a3, 0))
    connector(slide, 0, 0, 1, 1, begin=E(a3, 2), end=E(e3, 0), head=True)
    loop_x = cx[5] + cols[5] / 2 + 0.07
    edge_label(slide, loop_x, lane_y0 + 2 * lane_h - 0.09, 0.62, 0.17, p3_loop, align="l", bg=None)
    connector(slide, 0, 0, 1, 1, kind="hv", begin=E(e3, 3), end=E(diamonds[2], 2))
    connector(slide, 0, 0, 1, 1, kind="elbow", begin=E(diamonds[2], 3), end=E(r4, 1))
    e4_cx = (e4.left + e4.width / 2) / 914400
    connector(slide, e4_cx, r4.top / 914400, e4_cx, (e4.top + e4.height) / 914400, color=PURPLE)
    for label in [t_in, *p1, *p2, *p3, *p4]:
        if plain(label) not in plain(sec.mermaid()):
            warn(f"slide 6: diagram label drifted from content: {label}")


def s07_wireframes(slide, sec: Section, ctx: Ctx, deck: Deck):
    lines = sec.copy_lines()
    headline = lines[0]
    bullets = [ln[2:] for ln in lines[1:] if ln.startswith("- ")]
    text(slide, L, TOP - 0.02, 4.3, 0.36, Para(headline, weight=500), ctx, size=15, color=INK, name="headline",
         anchor="m")
    text(slide, L, TOP + 0.36, 4.3, 0.9, [Para(b, bullet=True, after=1.5) for b in bullets], ctx, size=10.5,
         color=INK2, name="trust bullets")

    # Workspace wireframe (spec §12), drawn natively.
    wx, wy, ww, wh = L, 2.8, 4.3, FOOT_Y - 0.1 - 2.8
    kicker(slide, wx, wy - 0.24, ww, "Workspace wireframe · spec §12")
    box(slide, wx, wy, ww, wh, fill=SURF, line=LINE, radius=0.06, name="wireframe frame")
    hb = box(slide, wx + 0.08, wy + 0.08, ww - 0.16, 0.26, fill=WHITE, line=LINE2, radius=0.04,
             name="wf header")
    text(slide, wx + 0.16, wy + 0.08, 1.2, 0.26, Para("Project", weight=600), size=7, color=INK, anchor="m",
         name="wf project")
    seg = box(slide, wx + 2.06, wy + 0.12, 0.84, 0.18, fill=SURF2, radius=0.09, name="wf ask-auto")
    box_text(seg, Para("Ask  |  Auto", align="c"), size=7, color=INK2, anchor="m")
    btn = box(slide, wx + 3.0, wy + 0.12, 1.14, 0.18, fill=INK, radius=0.09, name="wf sign-off")
    box_text(btn, Para("Request sign-off", align="c", weight=500), size=7, color=WHITE, anchor="m")
    body_y, body_h = wy + 0.42, wh - 0.42 - 0.42
    # agent lanes
    box(slide, wx + 0.08, body_y, 0.86, body_h, fill=WHITE, line=LINE2, radius=0.04, name="wf lanes")
    text(slide, wx + 0.13, body_y + 0.04, 0.8, 0.18, Para("Agent lanes", weight=600), size=7, color=INK,
         name="wf lanes label")
    for k, lane in enumerate(["Intake", "Layout", "Power", "Cooling"]):
        ly = body_y + 0.27 + k * 0.29
        box(slide, wx + 0.13, ly, 0.76, 0.22, fill=SURF2, radius=0.04, name="wf lane")
        text(slide, wx + 0.18, ly, 0.6, 0.22, Para(lane), size=7, color=INK2, anchor="m", name="wf lane label")
    # canvas
    cvx, cvw = wx + 1.02, ww - 1.02 - 0.08 - 0.96
    box(slide, cvx, body_y, cvw, body_h, fill=WHITE, line=LINE2, radius=0.04, name="wf canvas")
    tabs = ["3D", "SLD", "Matrix", "Options", "Docs"]
    tx = cvx + 0.06
    for k, tab in enumerate(tabs):
        tw = text_width(tab, 6.5, 500) + 0.14
        t = box(slide, tx, body_y + 0.05, tw, 0.17, fill=INK if k == 0 else SURF2, radius=0.08, name="wf tab")
        box_text(t, Para(tab, align="c", weight=500), size=7, color=WHITE if k == 0 else INK2, anchor="m")
        tx += tw + 0.04
    gx0, gy0, gw, gh = cvx + 0.12, body_y + 0.32, cvw - 0.24, body_h - 0.42
    for r_ in range(3):
        for c_ in range(8):
            box(slide, gx0 + c_ * gw / 8 + 0.02, gy0 + r_ * gh / 3 + 0.05, gw / 8 - 0.04, gh / 3 - 0.14,
                line=LINE, lw=0.5, name="wf rack")
    # co-pilot
    cpx = wx + ww - 0.08 - 0.9
    box(slide, cpx, body_y, 0.9, body_h, fill=WHITE, line=LINE2, radius=0.04, name="wf copilot")
    text(slide, cpx + 0.06, body_y + 0.04, 0.8, 0.18, Para("Co-pilot", weight=600), size=7, color=INK,
         name="wf copilot label")
    for k, (bw, al) in enumerate([(0.62, "r"), (0.7, "l"), (0.5, "r")]):
        bx = cpx + 0.06 if al == "l" else cpx + 0.84 - bw
        box(slide, bx, body_y + 0.3 + k * 0.27, bw, 0.18, fill=SURF2 if al == "l" else BLUE_T, radius=0.06,
            name="wf bubble")
    box(slide, cpx + 0.33, body_y + body_h - 0.3, 0.24, 0.24, fill=INK, shape=MSO_SHAPE.OVAL, name="wf voice")
    box(slide, cpx + 0.42, body_y + body_h - 0.21, 0.06, 0.06, fill=WHITE, shape=MSO_SHAPE.OVAL, name="wf voice dot")
    # revision timeline
    ty = wy + wh - 0.34
    box(slide, wx + 0.08, ty, ww - 0.16, 0.26, fill=WHITE, line=LINE2, radius=0.04, name="wf timeline")
    box(slide, wx + 0.4, ty + 0.125, ww - 0.9, 0.01, fill=LINE, name="wf timeline rule")
    for k, lab in enumerate(["r1", "r2", "signed"]):
        px = wx + 0.4 + k * (ww - 0.9) / 2
        box(slide, px - 0.04, ty + 0.09, 0.08, 0.08, fill=INK if k == 2 else MUTED, shape=MSO_SHAPE.OVAL,
            name="wf rev")
        text(slide, px + 0.07, ty, 0.6, 0.26, Para(lab, mono=True), size=6.5, color=INK2, anchor="m",
             name="wf rev label")

    # Annotated finding drawer: numbered legend beside the screenshot.
    callout_m = re.search(r"callouts on (.+?)\.\s*$", sec.visual, re.M)
    callouts = [c.strip().strip(",").strip() for c in re.split(r"\(\d\)", callout_m.group(1)) if c.strip()]
    callouts = [c.removesuffix(" and").strip().strip(",").strip() for c in callouts]
    if len(callouts) != 6:
        raise SystemExit("slide 7: expected six callouts")
    rx = 4.95
    kicker(slide, rx, TOP - 0.02, 2.0, "Finding drawer · annotated")
    pitch = (FOOT_Y - 0.1 - TOP - 0.3) / 6
    for k, c in enumerate(callouts):
        y = TOP + 0.3 + k * pitch
        number_badge(slide, str(k + 1), rx, y + 0.02)
        text(slide, rx + 0.32, y - 0.02, 1.68, 0.46, Para(c[0].upper() + c[1:]), ctx, size=9.5, color=INK2,
             name="callout")
    deck.shots.place(slide, 7, 0, 7.05, TOP, R - 7.05, FOOT_Y - 0.1 - TOP, valign="t")
    footer(slide, sec.quoted("Footer (small print)"), ctx)


def s08_architecture(slide, sec: Section, ctx: Ctx, deck: Deck):
    paras = [p.strip() for p in sec.copy.split("\n\n") if p.strip()]
    text(slide, L, TOP - 0.04, 4.0, 0.34, Para(paras[0], weight=500), ctx, size=15, color=INK, name="headline",
         anchor="m")
    text(slide, L, TOP + 0.3, CW, 0.4, Para(paras[1]), ctx, size=9.5, color=INK2, name="subline")
    src = sec.mermaid()
    n = _mermaid_nodes(src)
    edges = _mermaid_edge_labels(src)
    sub = re.findall(r'subgraph (\w+)\["(.+?)"\]', src)
    groups = dict(sub)

    def split(label: str) -> tuple[str, str]:
        return tuple(s.strip() for s in label.split(": ", 1)) if ": " in label else (label, "")

    ax, aw = L, 1.78
    bx, bw = 2.97, 3.27
    cx_, cw = 7.17, R - 7.17
    top = 2.0
    # Column A: browser
    browser = box(slide, ax, top, aw, 1.5, fill=SURF, line=LINE, radius=0.06, name="browser group")
    bt, bs = split(groups["B"])
    text(slide, ax + 0.1, top + 0.05, aw - 0.2, 0.34, [Para(bt, weight=600, size=8.5), Para(bs, size=7.5, color=MUTED)],
         size=8, color=INK, name="group label", lnspc=0.95)
    ut, us = split(n["UI"])
    ui = diagram_card(slide, ax + 0.1, top + 0.42, aw - 0.2, 0.56, ut, us, name="workspace ui")
    diagram_card(slide, ax + 0.1, top + 1.06, aw - 0.2, 0.34, n["SD"], name="demo world", title_weight=500)
    # Firebase services and Gemini Live below the browser
    fb_y = top + 1.78
    fb = box(slide, ax, fb_y, aw, 5.27 - fb_y, fill=SURF, line=LINE, radius=0.06, name="firebase group")
    yy = fb_y + 0.1
    for key in ("AUTH", "AC", "LIVE"):
        diagram_card(slide, ax + 0.1, yy, aw - 0.2, 0.32 if key != "LIVE" else 0.4, n[key], name=key.lower(),
                     title_weight=500, size=8)
        yy += (0.32 if key != "LIVE" else 0.4) + 0.07
    voice = next(e for e in edges if e.startswith("voice"))
    text(slide, ax + 0.1, yy - 0.04, aw - 0.2, 0.18, Para(voice, align="c"), size=7, color=MUTED,
         name="voice label")
    connector(slide, ax + aw / 2, top + 1.5, ax + aw / 2, fb_y, name="browser to firebase")

    # Column B: Cloud Run kirchhoff-api
    pb_h = 2.42
    panel = box(slide, bx, top, bw, pb_h, fill=SURF, line=LINE, radius=0.06, name="kirchhoff-api group")
    gt = groups["CR1"]
    text(slide, bx + 0.12, top + 0.04, bw - 0.24, 0.22, Para(gt, weight=600), size=8, color=INK, anchor="m",
         name="group label")
    at, as_ = split(n["API"])
    api = diagram_card(slide, bx + 0.12, top + 0.3, bw - 0.24, 0.34, at, as_, name="fastapi", align="c")
    kt, ks = split(n["ADK"])
    adk = diagram_card(slide, bx + 0.12, top + 0.78, bw - 0.24, 0.6, kt, ks, fill=BLUE_T, line=BLUE, color=INK,
                       name="adk agents")
    zone_y = top + 1.5
    zw = bw - 0.12 - 0.3                  # inset from the panel edge so deploy arrows land on the service
    zone = box(slide, bx + 0.12, zone_y, zw, 0.82, line=RED, lw=1.25, dash=MSO_LINE.DASH, radius=0.05,
               name="no-LLM zone")
    text(slide, bx + 0.2, zone_y + 0.03, zw - 0.16, 0.18, Para(groups["DET"], weight=600, align="r"), size=7.5,
         color=RED, anchor="m", name="zone label")
    et, es = split(n["ENG"])
    rt, rs = split(n["RP"])
    ew = (zw - 0.2 - 0.3) / 2
    eng = diagram_card(slide, bx + 0.22, zone_y + 0.25, ew, 0.48, et, es, name="engine")
    rp = diagram_card(slide, bx + 0.22 + ew + 0.3, zone_y + 0.25, ew, 0.48, rt, rs, name="checker")
    connector(slide, 0, 0, 1, 1, begin=(api, 2), end=(adk, 0), name="api to adk")
    connector(slide, bx + 0.22 + ew / 2, top + 1.38, bx + 0.22 + ew / 2, zone_y + 0.25, name="adk to engine")
    connector(slide, 0, 0, 1, 1, begin=(eng, 3), end=(rp, 1), name="engine to checker")

    # Below the panel: Firestore and the private enterprise service.
    r4 = top + pb_h + 0.32
    fs_w = 1.42
    fs = box(slide, bx, r4, fs_w, 5.27 - r4, fill=WHITE, line=INK2, shape=MSO_SHAPE.CAN, name="firestore")
    box_text(fs, Para(n["FS"], weight=600, align="c"), size=8, anchor="m", margin=(0.03, 0.12, 0.03, 0.02))
    ent_x = bx + fs_w + 0.36
    en_t, en_s = n["ENT"].split(", ", 1) if ", " in n["ENT"] else (n["ENT"], "")
    ent = diagram_card(slide, ent_x, r4, bx + bw - ent_x, 5.27 - r4, en_t, en_s, fill=PURPLE_T, line=PURPLE,
                       name="kirchhoff-enterprise", size=7.5, sub_size=7)
    lab = {e.split(",")[0]: e for e in edges}
    connector(slide, bx + fs_w / 2, top + pb_h, bx + fs_w / 2, r4 + 0.06, name="api to firestore")
    text(slide, bx + fs_w / 2 + 0.07, top + pb_h + 0.02, 1.1, 0.28, Para(lab["runs"]), size=7, color=MUTED,
         anchor="m", name="edge label")
    ent_cx = ent_x + (bx + bw - ent_x) / 2
    connector(slide, ent_cx, top + pb_h, ent_cx, r4, name="api to enterprise")
    text(slide, ent_cx + 0.07, top + pb_h + 0.02, 0.95, 0.28, Para(lab["Google ID token"]), size=7, color=MUTED,
         anchor="m", name="edge label")
    connector(slide, ent_x, r4 + 0.5, bx + fs_w, r4 + 0.5, name="enterprise to firestore")

    # Browser <-> API, Firestore -> browser (read-only)
    y_api = top + 0.3 + 0.17
    connector(slide, ax + aw, y_api, bx + 0.12, y_api, name="ui to api")
    edge_label(slide, ax + aw + 0.03, y_api - 0.2, bx - ax - aw - 0.06, 0.16, lab["/api rewrite"], bg=None)
    y_back = top + 0.88
    connector(slide, bx, r4 + 0.36, ax + aw, y_back, kind="elbow", name="firestore to ui")
    mid = (ax + aw + bx) / 2
    edge_label(slide, mid - 0.37, 3.5, 0.74, 0.3, lab["onSnapshot"])

    # Column C: Cloud Storage, Gemini, Cloud Build
    gcs = box(slide, cx_ + 0.2, top + 0.24, cw - 0.4, 0.46, fill=WHITE, line=INK2, shape=MSO_SHAPE.CAN,
              name="cloud storage")
    box_text(gcs, Para(n["GCS"], weight=600, align="c"), size=8, anchor="m", margin=(0.03, 0.1, 0.03, 0.02))
    connector(slide, bx + bw - 0.12, y_api, cx_ + 0.2, y_api, name="api to storage")
    edge_label(slide, bx + bw + 0.02, y_api - 0.2, cx_ + 0.2 - bx - bw - 0.04, 0.16, lab["signed URLs"], bg=None)
    gm = n["GEM"]
    gt_, gs_ = (gm.split(" on ", 1)[0], "on " + gm.split(" on ", 1)[1]) if " on " in gm else (gm, "")
    gem = diagram_card(slide, cx_, top + 0.78, cw, 0.6, gt_, gs_, fill=BLUE_T, line=BLUE, name="gemini", size=8.5)
    y_gem = top + 1.08
    connector(slide, bx + bw - 0.12, y_gem, cx_, y_gem, name="adk to gemini")
    edge_label(slide, bx + bw + 0.02, y_gem - 0.36, cx_ - bx - bw - 0.04, 0.33, lab["structured output"], bg=None)
    cb_y = top + 1.62
    ct, cs = split(n["CB"])
    cb = diagram_card(slide, cx_, cb_y, cw, 0.42, ct, cs, name="cloud build", title_weight=600)
    y_cb = cb_y + 0.21
    connector(slide, cx_, y_cb, bx + bw, y_cb, dash=MSO_LINE.DASH, name="build to api")
    connector(slide, cx_, y_cb, bx + bw, r4 + 0.24, kind="elbow", dash=MSO_LINE.DASH, name="build to enterprise")
    for key in ("UI", "API", "ADK", "ENG", "RP", "GEM", "FS", "GCS", "ENT", "CB", "AUTH", "AC", "LIVE", "SD"):
        if key not in n:
            raise SystemExit(f"slide 8: mermaid node {key} missing")


def s09_tech(slide, sec: Section, ctx: Ctx, deck: Deck):
    rows = []
    for ln in sec.copy_lines():
        m = re.match(r"^\*\*(.+?):\*\*\s*(.+)$", ln)
        if m:
            rows.append((m.group(1), [s.strip() for s in m.group(2).split(" · ")]))
    if len(rows) != 5:
        raise SystemExit("slide 9: expected five technology rows")
    y = TOP + 0.12
    pitch = 0.8
    for i, (label, items) in enumerate(rows):
        if i:
            rule(slide, L, y - 0.19, R, y - 0.19, color=LINE2)
        kicker(slide, L, y + 0.11, 1.6, label, color=INK2, size=8)
        x = L + 1.75
        for item in items:
            size = 11
            w = text_width(item, size, 500) + 0.36
            chip = box(slide, x, y, w, 0.42, fill=WHITE, line=LINE, radius=0.21, name="tech chip")
            box_text(chip, Para(item, weight=500, align="c"), size=size, color=INK, anchor="m",
                     margin=(0.06, 0, 0.06, 0), min_scale=0.9)
            x += w + 0.14
            if x > R + 0.01:
                warn("slide 9: chip row overflows the slide")
        y += pitch


def s10_cost(slide, sec: Section, ctx: Ctx, deck: Deck):
    blocks = [b.strip() for b in sec.copy.split("\n\n") if b.strip()]
    run_lines = blocks[0].splitlines()
    run_title = plain(run_lines[0])
    run_items = [ln[2:] for ln in run_lines[1:] if ln.startswith("- ")]
    charge_lines = blocks[1].splitlines()
    charge_title, charge_text = plain(charge_lines[0]), " ".join(charge_lines[1:])
    ch = 2.05
    card(slide, L, TOP, 4.7, ch, name="run cost card")
    icon_badge(slide, "receipt", L + 0.2, TOP + 0.18, 0.38, bg=WHITE)
    text(slide, L + 0.7, TOP + 0.18, 3.6, 0.38, Para(run_title, weight=600), size=14, color=INK, anchor="m",
         name="card title")
    text(slide, L + 0.22, TOP + 0.74, 4.3, ch - 0.86, [Para(it, bullet=True, after=4) for it in run_items], ctx,
         size=12, color=INK2, name="run cost bullets")
    x2 = L + 4.7 + 0.2
    card(slide, x2, TOP, R - x2, ch, name="charge card")
    icon_badge(slide, "users", x2 + 0.2, TOP + 0.18, 0.38, bg=WHITE)
    text(slide, x2 + 0.7, TOP + 0.18, 3.3, 0.38, Para(charge_title, weight=600), size=14, color=INK, anchor="m",
         name="card title")
    text(slide, x2 + 0.22, TOP + 0.74, R - x2 - 0.44, 1.2, Para(charge_text, weight=500), ctx, size=15,
         color=INK, name="charge text")

    # Integration path strip.
    strip = re.search(r"integration path: (.+?), labelled \"(.+?)\"", sec.visual)
    path, caption = strip.group(1), strip.group(2)
    stages = [plain(s).strip() for s in path.split(" → ")]
    erp, crm = [s.strip() for s in stages[2].split(" / ")]
    sy = TOP + ch + 0.36
    kicker(slide, L, sy, 3.0, "Enterprise integration")
    by_ = sy + 0.32
    bh = 0.4
    k = diagram_card(slide, L, by_ + 0.22, 1.5, bh, stages[0], fill=INK, line=INK, color=WHITE, size=10,
                     name="strip kirchhoff")
    api_w = text_width(stages[1], 9.5, 500, mono=True) + 0.4
    a = box(slide, L + 1.5 + 0.55, by_ + 0.22, api_w, bh, fill=WHITE, line=LINE, radius=0.05, name="strip api")
    box_text(a, Para(stages[1], mono=True, weight=500, align="c"), size=9.5, color=INK, anchor="m")
    ex = L + 1.5 + 0.55 + api_w + 0.6
    e1 = diagram_card(slide, ex, by_, 2.1, 0.36, erp, fill=PURPLE_T, line=PURPLE, size=9.5, title_weight=500,
                      name="strip erp")
    e2 = diagram_card(slide, ex, by_ + 0.48, 2.1, 0.36, crm, fill=PURPLE_T, line=PURPLE, size=9.5,
                      title_weight=500, name="strip crm")
    connector(slide, 0, 0, 1, 1, begin=(k, 3), end=(a, 1))
    connector(slide, 0, 0, 1, 1, kind="elbow", begin=(a, 3), end=(e1, 1))
    connector(slide, 0, 0, 1, 1, kind="elbow", begin=(a, 3), end=(e2, 1))
    text(slide, ex + 2.3, by_ + 0.12, R - ex - 2.3, 0.6, Para(caption), ctx, size=9.5, color=MUTED, anchor="m",
         name="strip caption")
    # Footer: source for the superscript on the slide.
    sups = {"¹²": 12}
    notes = []
    for mark, num in sups.items():
        if mark in sec.copy and num in deck.sources:
            notes.append(f"{mark} " + re.sub(r"\s*https?://\S+$", "", deck.sources[num]))
    if notes:
        footer(slide, "  ".join(notes), ctx)


def s11_snapshots(slide, sec: Section, ctx: Ctx, deck: Deck):
    caps = [re.sub(r"^\d+\.\s*", "", ln) for ln in sec.copy_lines() if re.match(r"^\d+\.", ln)]
    if len(caps) != 6:
        raise SystemExit("slide 11: expected six captions")
    gap_x, cap_h, gap_y = 0.2, 0.26, 0.08
    w = (CW - 2 * gap_x) / 3
    h = min(w * 9 / 16, (FOOT_Y - 0.06 - TOP - 2 * cap_h - gap_y) / 2)
    for i, c in enumerate(caps):
        x = L + (i % 3) * (w + gap_x)
        y = TOP + (i // 3) * (h + cap_h + gap_y)
        deck.shots.place(slide, 11, i, x, y, w, h)
        number_badge(slide, str(i + 1), x, y + h + 0.05, d=0.18, size=7)
        text(slide, x + 0.25, y + h + 0.02, w - 0.25, 0.24, Para(c, weight=500), ctx, size=9.5, color=INK,
             anchor="m", name="snapshot caption")
    footer(slide, sec.quoted("Footer (small print)"), ctx)


def s12_performance(slide, sec: Section, ctx: Ctx, deck: Deck):
    bullets = [ln[2:] for ln in sec.copy_lines() if ln.startswith("- ")]
    stats, estimate = [], None
    for b in bullets:
        m = re.match(r"^\*\*(.+?)\*\*\s*(.*)$", b)
        if m.group(1).rstrip(":").lower() == "estimate":
            estimate = m.group(2)
        else:
            stats.append((m.group(1), m.group(2)))
    if len(stats) != 4 or not estimate:
        raise SystemExit("slide 12: expected four stats and one estimate")
    icons = ["shield-check", "list-checks", "timer", "fingerprint"]
    widths = [2.05, 2.05, 2.05, CW - 3 * 2.05 - 3 * 0.15]
    x = L
    th = 1.3
    for (big, label), ic, w in zip(stats, icons, widths):
        card(slide, x, TOP, w, th, name="stat tile")
        icon(slide, ic, x + w - 0.36, TOP + 0.14, 0.22)
        text(slide, x + 0.16, TOP + 0.1, w - 0.55, 0.52, Para(big, weight=500), ctx, size=24, color=INK,
             anchor="m", name="stat value", min_scale=0.8)
        text(slide, x + 0.16, TOP + 0.66, w - 0.3, th - 0.72, Para(label), ctx, size=9.5, color=INK2,
             name="stat label")
        x += w + 0.15
    sy = TOP + th + 0.18
    sh = FOOT_Y - 0.1 - sy
    sw = sh * 16 / 9                       # the shot slot keeps the capture's 16:9 shape
    deck.shots.place(slide, 12, 0, L, sy, sw, sh)
    ex = L + sw + 0.2
    ew = R - ex
    box(slide, ex, sy, ew, sh, fill=WHITE, line=LINE, lw=0.75, dash=MSO_LINE.DASH, radius=0.08,
        name="estimate tile")
    tag = box(slide, ex + 0.2, sy + 0.2, 0.86, 0.22, fill=SURF2, radius=0.11, name="estimate tag")
    box_text(tag, Para("Estimate", caps=True, mono=True, spc=0.6, align="c", weight=500), size=7, color=MUTED,
             anchor="m")
    icon(slide, "clock", R - 0.44, sy + 0.2, 0.22)
    m = re.match(r"^(~?\d+ engineer-hours)\s*(.*)$", estimate)
    if m:
        text(slide, ex + 0.2, sy + 0.62, ew - 0.4, 0.5, Para(m.group(1), weight=500), ctx, size=24, color=INK,
             name="estimate value", anchor="m", min_scale=0.8)
        text(slide, ex + 0.2, sy + 1.16, ew - 0.4, sh - 1.3, Para(m.group(2)), ctx, size=11, color=INK2,
             name="estimate label")
    else:
        text(slide, ex + 0.2, sy + 0.62, ew - 0.4, sh - 0.8, Para(estimate), ctx, size=13, color=INK,
             name="estimate text")
    footer(slide, sec.quoted("Footer (small print)"), ctx)


def s13_future(slide, sec: Section, ctx: Ctx, deck: Deck):
    cols = []
    for ln in sec.copy_lines():
        m = re.match(r"^\*\*(.+?):\*\*\s*(.+)$", ln)
        if m:
            cols.append((m.group(1), m.group(2)))
    if len(cols) != 3:
        raise SystemExit("slide 13: expected Now / Next / Smart engineering")
    icons = ["map-pin", "rocket", "network"]
    gap = 0.2
    w = (CW - 2 * gap) / 3
    ch = 1.2
    for i, ((label, body), ic) in enumerate(zip(cols, icons)):
        x = L + i * (w + gap)
        card(slide, x, TOP, w, ch, name="roadmap column")
        icon(slide, ic, x + 0.16, TOP + 0.15, 0.2)
        text(slide, x + 0.44, TOP + 0.12, w - 0.6, 0.26, Para(label, weight=600), size=11, color=INK,
             anchor="m", name="roadmap label")
        text(slide, x + 0.16, TOP + 0.44, w - 0.32, ch - 0.5, Para(body[0].upper() + body[1:]), ctx, size=10.5,
             color=INK2, name="roadmap text")
    sy = TOP + ch + 0.18
    sw = (CW - gap) / 2
    sh = min(sw * 9 / 16, FOOT_Y - 0.1 - sy)
    deck.shots.place(slide, 13, 0, L, sy, sw, sh)
    deck.shots.place(slide, 13, 1, L + sw + gap, sy, sw, sh)
    footer(slide, sec.quoted("Footer (small print)"), ctx)


def s14_links(slide, sec: Section, ctx: Ctx, deck: Deck):
    keep_heading_only(slide, "Provide links to your:")
    rows = []
    for ln in sec.copy_lines():
        m = re.match(r"^- (?P<label>[^:]+): (?P<val>\[TBD: [^\]]+\]|\S+)\s*(?P<note>.*)$", ln)
        if m:
            rows.append((m.group("label"), m.group("val"), m.group("note")))
    if len(rows) != 3:
        raise SystemExit("slide 14: expected three link rows")
    icons = ["github", "circle-play", "globe"]
    rh, gap = 1.18, 0.12
    for i, ((label, val, note), ic) in enumerate(zip(rows, icons)):
        y = TOP + i * (rh + gap)
        card(slide, L, y, CW, rh, name="link row")
        icon_badge(slide, ic, L + 0.22, y + (rh - 0.5) / 2, 0.5, bg=WHITE)
        text(slide, L + 0.95, y + 0.16, 6.6, 0.24, Para(label, weight=500), size=10.5, color=MUTED,
             name="template prompt")
        runs = ctx.runs(val)
        url = runs[0].text if runs and runs[0].kind == "text" and runs[0].text.startswith("http") else None
        if url:
            runs = [Run(url, mono=True, weight=500, link=url)]
        else:
            for r in runs:
                r.mono = r.kind != "tbd"
        text(slide, L + 0.95, y + 0.42, 6.75, 0.36, Para(runs), size=15, color=INK, anchor="m", name="link url",
             min_scale=0.75)
        if note:
            text(slide, L + 0.95, y + 0.8, 6.75, 0.24, Para(note), ctx, size=9.5, color=MUTED, name="link note")
        qx, qs = R - 0.2 - 0.9, 0.9
        if url:
            pic = slide.shapes.add_picture(qr_png(url), I(qx), I(y + (rh - qs) / 2), I(qs), I(qs))
            pic.name = f"QR {label}"
        else:
            q = box(slide, qx, y + (rh - qs) / 2, qs, qs, fill=WHITE, line=SHOT_LINE, lw=1.0, dash=MSO_LINE.DASH,
                    radius=0.06, name="QR placeholder")
            box_text(q, [Para(t, align="c") for t in balance("QR appears when the URL is filled", 0.62, 6.5)],
                     size=6.5, color=SHOT_TEXT,
                     anchor="m", margin=(0.06, 0.04, 0.06, 0.04))


def s15_team(slide, sec: Section, ctx: Ctx, deck: Deck):
    lines = sec.copy_lines()
    heading = plain(lines[0])
    add_heading(slide, heading)
    people = [ln[2:] for ln in lines if ln.startswith("- ")]
    closing = [ln for ln in lines[1:] if not ln.startswith("- ")][0]
    photo_keys = ["photos.mehul", "photos.jayadratha"]
    w, gap = (CW - 0.3) / 2, 0.3
    ch = 2.1
    for i, person in enumerate(people):
        m = re.match(r"^(?P<name>[^,.]+)[,.]\s*(?P<rest>.*)$", person)
        name, rest = m.group("name").strip(), m.group("rest").strip()
        x = L + i * (w + gap)
        card(slide, x, TOP, w, ch, name="person card")
        d = 1.5
        px, py = x + 0.25, TOP + (ch - d) / 2
        photo = deck.fill.get(photo_keys[i])
        ppath = (deck.fill.base / photo) if photo else None
        if ppath and ppath.exists():
            pic = picture(slide, ppath, px, py, d, d, mode="cover", border=False, name=f"photo {name}")
            pic.auto_shape_type = MSO_SHAPE.OVAL
        else:
            deck.fill.open.append(OpenItem("PHOTO", 15, photo_keys[i], f"headshot of {name}"))
            box(slide, px, py, d, d, fill=SHOT_FILL, line=SHOT_LINE, lw=1.0, dash=MSO_LINE.DASH,
                shape=MSO_SHAPE.OVAL, name="photo placeholder")
            icon(slide, "user", px + d / 2 - 0.25, py + 0.28, 0.5, "muted")
            text(slide, px + 0.15, py + 0.86, d - 0.3, 0.4, Para(f"PHOTO · {photo_keys[i]}", align="c", mono=True),
                 size=6.5, color=SHOT_TEXT, name="photo label")
        tx = px + d + 0.3
        text(slide, tx, TOP + 0.42, x + w - tx - 0.2, 0.42, Para(name, weight=600), size=18, color=INK,
             anchor="m", name="person name")
        text(slide, tx, TOP + 0.9, x + w - tx - 0.2, 1.0, Para(rest[0].upper() + rest[1:] if rest[0] != "[" else rest),
             ctx, size=11, color=INK2, name="person role")
    text(slide, L, TOP + ch + 0.35, CW, 0.4, Para(closing, weight=500), ctx, size=16, color=INK, name="closing")
    theme = re.search(r'small strip: "(.+?)"', sec.visual).group(1)
    label, value = theme.split(": ", 1)
    sy = TOP + ch + 0.95
    strip = box(slide, L, sy, CW, 0.5, fill=SURF, line=LINE2, radius=0.08, name="theme strip")
    text(slide, L + 0.22, sy, 1.0, 0.5, Para(label, caps=True, mono=True, spc=0.8, weight=500), size=7.5,
         color=MUTED, anchor="m", name="theme label")
    text(slide, L + 1.2, sy, CW - 1.4, 0.5, Para(value, weight=500), size=13, color=INK, anchor="m",
         name="theme value")


def s16_thanks(slide, sec: Section, ctx: Ctx, deck: Deck):
    lines = sec.copy_lines()
    text(slide, 5.3, 2.42, 4.3, 0.4, Para(lines[0], weight=600), size=18, color=WHITE, name="closing product")
    text(slide, 5.3, 2.84, 4.3, 0.34, Para(lines[1]), size=12.5, color="E8EAED", name="closing tagline")
    url = lines[2] if lines[2].startswith("http") else "https://" + lines[2]
    text(slide, 5.3, 3.22, 4.3, 0.3, Para([Run(lines[2], mono=True, link=url)]), size=12, color="E8EAED",
         name="closing url")


BUILDERS = {
    1: s01_team, 3: s03_brief, 4: s04_opportunities, 5: s05_features, 6: s06_process, 7: s07_wireframes,
    8: s08_architecture, 9: s09_tech, 10: s10_cost, 11: s11_snapshots, 12: s12_performance, 13: s13_future,
    14: s14_links, 15: s15_team, 16: s16_thanks,
}


# --------------------------------------------------------------------------------------------
# Build.
# --------------------------------------------------------------------------------------------
def write_deterministic(src: Path, dst: Path) -> None:
    with zipfile.ZipFile(src) as zin, zipfile.ZipFile(dst, "w", zipfile.ZIP_DEFLATED) as zout:
        for info in zin.infolist():
            data = zin.read(info.filename)
            zi = zipfile.ZipInfo(info.filename, date_time=(2026, 10, 17, 0, 0, 0))
            zi.compress_type = zipfile.ZIP_DEFLATED
            zi.external_attr = 0o644 << 16
            zout.writestr(zi, data)


def export_pdf(pptx: Path) -> Path | None:
    exe = shutil.which("soffice") or shutil.which("libreoffice")
    if not exe:
        print("PDF: LibreOffice (soffice) is not installed here; no PDF written.")
        return None
    subprocess.run([exe, "--headless", "--convert-to", "pdf", "--outdir", str(pptx.parent), str(pptx)], check=True)
    pdf = pptx.with_suffix(".pdf")
    print(f"PDF: {pdf}")
    return pdf


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--template", type=Path, default=TEMPLATE)
    ap.add_argument("--content", type=Path, default=CONTENT_MD)
    ap.add_argument("--out", type=Path, default=OUT_PPTX)
    ap.add_argument("--shots", type=Path, default=SHOTS_DIR, help="screenshot folder (default: shots/)")
    ap.add_argument("--fill", type=Path, default=FILL_YAML, help="team-fill values (default: fill.yaml)")
    ap.add_argument("--pdf", action="store_true", help="also export a PDF with LibreOffice, if installed")
    ap.add_argument("--strict", action="store_true", help="exit 1 while any SHOT, TBD, PHOTO or (spec) slot is open")
    args = ap.parse_args()

    PATHS["shots"] = args.shots.resolve()
    sections, sources = parse_content(args.content.read_text(encoding="utf-8"))
    fill = Fill(args.fill.resolve())
    shots = Shots(sections, fill)
    tag_line = sections[16].copy_lines()
    product, by = tag_line[0].split(" by ", 1)
    deck = Deck(sections, sources, fill, shots, tagline=tag_line[1], product=product, by=f"by {by}")

    prs = Presentation(str(args.template))
    slides = {i + 1: s for i, s in enumerate(prs.slides)}
    if len(slides) != 16:
        raise SystemExit(f"template: expected 16 slides, found {len(slides)}")
    for n, build in BUILDERS.items():
        ctx = Ctx(fill, n)
        build(slides[n], sections[n], ctx, deck)
        set_notes(slides[n], ctx.notes(sections[n].notes))
    fill.check_counts()
    delete_slide(prs, 1)        # template slide 2: download instructions

    cp = prs.core_properties
    cp.title = "Kirchhoff by Ordinal: AI Builder Cup 2026"
    cp.author = "Team Ordinal"
    cp.last_modified_by = "Team Ordinal"
    cp.subject = "AI Builder Cup 2026 submission"
    cp.keywords = ""
    cp.comments = ""
    cp.revision = 1
    cp.created = cp.modified = datetime(2026, 10, 17, 0, 0, 0)     # fixed, so rebuilds are byte-identical

    tmp = args.out.with_suffix(".tmp.pptx")
    prs.save(str(tmp))
    write_deterministic(tmp, args.out)
    tmp.unlink()
    PATHS["shots"].mkdir(exist_ok=True)
    (PATHS["shots"] / "README.md").write_text(shots.readme(), encoding="utf-8")

    print(f"Built {args.out} ({len(prs.slides)} slides)")
    if fill.open:
        print(f"\nOpen slots ({len(fill.open)}), numbered by template slide:")
        for it in fill.open:
            print(f"  [{it.kind:5}] slide {it.slide:>2}  {it.key}  {it.text[:90]}")
    if WARNINGS:
        print(f"\nLayout warnings ({len(WARNINGS)}):")
        for w in WARNINGS:
            print(f"  - {w}")
    if args.pdf:
        export_pdf(args.out)
    if args.strict and fill.open:
        print("\n--strict: open slots remain")
        return 1
    return 0


if __name__ == "__main__":
    sys.exit(main())
